from pathlib import Path
import json
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path(__file__).resolve().parent / 'output'
out.mkdir(exist_ok=True)
doc = Document()
doc.styles['Normal'].font.name = 'Times New Roman'
doc.styles['Normal'].font.size = Pt(12)
doc.add_heading('MathType automation validation', 0)
doc.add_paragraph('Inline: [[EQ:inline]] remains within this sentence.')
doc.add_paragraph('Standalone inline object:')
doc.add_paragraph('[[EQ:inline-standalone]]')
bookmark_paragraph = doc.add_paragraph('Bookmark inline: ')
bookmark_start = OxmlElement('w:bookmarkStart')
bookmark_start.set(qn('w:id'), '42')
bookmark_start.set(qn('w:name'), 'mt_bookmark_inline')
bookmark_end = OxmlElement('w:bookmarkEnd')
bookmark_end.set(qn('w:id'), '42')
bookmark_paragraph._p.append(bookmark_start)
bookmark_paragraph._p.append(bookmark_end)
bookmark_paragraph.add_run(' inserted at an empty bookmark.')
doc.add_paragraph('Display equation:')
display_bookmark_paragraph = doc.add_paragraph()
display_bookmark_start = OxmlElement('w:bookmarkStart')
display_bookmark_start.set(qn('w:id'), '43')
display_bookmark_start.set(qn('w:name'), 'mt_bookmark_display')
display_bookmark_end = OxmlElement('w:bookmarkEnd')
display_bookmark_end.set(qn('w:id'), '43')
display_bookmark_paragraph._p.append(display_bookmark_start)
display_bookmark_paragraph.add_run('[[EQ:display]]')
display_bookmark_paragraph._p.append(display_bookmark_end)
doc.add_paragraph('Display equation addressed later by MathType index:')
doc.add_paragraph('[[EQ:display-index]]')
doc.add_paragraph('Native right-numbered equation:')
doc.add_paragraph('[[EQ:numbered]]')
doc.add_paragraph('Second native right-numbered equation:')
doc.add_paragraph('[[EQ:numbered2]]')
doc.add_paragraph('End of validation document.')
field_paragraph = doc.add_paragraph('Unrelated field sentinel: ')
run = field_paragraph.add_run()
begin = OxmlElement('w:fldChar')
begin.set(qn('w:fldCharType'), 'begin')
instruction = OxmlElement('w:instrText')
instruction.set(qn('xml:space'), 'preserve')
instruction.text = ' AUTHOR '
separate = OxmlElement('w:fldChar')
separate.set(qn('w:fldCharType'), 'separate')
text = OxmlElement('w:t')
text.text = 'UNCHANGED_FIELD'
end = OxmlElement('w:fldChar')
end.set(qn('w:fldCharType'), 'end')
for element in (begin, instruction, separate, text, end):
    run._r.append(element)
doc.save(out / 'input.docx')
config = {
    'input': 'input.docx', 'output': 'result.docx', 'pdf': 'result.pdf',
    'equations': [
        {'anchor': '[[EQ:inline]]', 'latex': r'E=mc^2', 'mode': 'inline'},
        {'anchor': '[[EQ:inline-standalone]]', 'latex': r'a^2+b^2=c^2', 'mode': 'inline'},
        {'bookmark': 'mt_bookmark_inline', 'latex': r'p=mv', 'mode': 'inline'},
        {'bookmark': 'mt_bookmark_display', 'latex': r'\int_0^1 x^2\,dx=\frac{1}{3}', 'mode': 'display'},
        {'anchor': '[[EQ:display-index]]', 'latex': r'\lim_{x\to 0}\frac{\sin x}{x}=1', 'mode': 'display'},
        {'anchor': '[[EQ:numbered]]', 'latex': r'\frac{-b\pm\sqrt{b^2-4ac}}{2a}', 'mode': 'right-numbered'},
        {'anchor': '[[EQ:numbered2]]', 'latex': r'\sum_{i=1}^{n}i=\frac{n(n+1)}{2}', 'mode': 'right-numbered'},
    ],
}
(out / 'equations.json').write_text(json.dumps(config, indent=2), encoding='utf-8')

number_existing = {
    'input': 'result.docx',
    'output': 'renumbered.docx',
    'pdf': 'renumbered.pdf',
    'equations': [
        {'bookmark': 'mt_bookmark_display', 'operation': 'number-existing'},
        {'equationIndex': 5, 'operation': 'number-existing'},
    ],
}
(out / 'number-existing.json').write_text(
    json.dumps(number_existing, indent=2), encoding='utf-8'
)
print(out)
