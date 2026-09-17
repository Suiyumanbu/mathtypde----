from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import statistics
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def prepare(root: Path, counts: list[int], repeats: int, filler_count: int) -> None:
    for count in counts:
        case = root / f'n-{count}'
        case.mkdir(parents=True)
        doc = Document()
        doc.styles['Normal'].font.size = Pt(11)
        seed_items = []
        recreate_items = []
        reuse_items = []
        for index in range(1, count + 1):
            for _ in range(filler_count // count + (index <= filler_count % count)):
                doc.add_paragraph('Performance fixture body text. ' * 8)
            anchor = f'[[PERF:{index:04d}]]'
            bookmark = f'mt_perf_{index:04d}'
            paragraph = doc.add_paragraph()
            begin = OxmlElement('w:bookmarkStart')
            begin.set(qn('w:id'), str(index))
            begin.set(qn('w:name'), bookmark)
            end = OxmlElement('w:bookmarkEnd')
            end.set(qn('w:id'), str(index))
            paragraph._p.append(begin)
            paragraph.add_run(anchor)
            paragraph._p.append(end)
            latex = rf'\frac{{a_{{{index}}}+b}}{{1+x^2}}=c_{{{index}}}'
            seed_items.append({'bookmark': bookmark, 'latex': latex, 'mode': 'display'})
            recreate_items.append({'bookmark': bookmark, 'latex': latex, 'mode': 'right-numbered'})
            reuse_items.append({'bookmark': bookmark, 'operation': 'number-existing'})
        doc.save(case / 'anchors.docx')
        seed = {'input': 'anchors.docx', 'output': 'display.docx', 'equations': seed_items[::-1]}
        (case / 'seed.json').write_text(json.dumps(seed, indent=2), encoding='utf-8')
        for repeat in range(1, repeats + 1):
            for strategy, source, items in (
                ('recreate', 'anchors.docx', recreate_items),
                ('reuse', 'display.docx', reuse_items),
            ):
                config = {'input': source, 'output': f'{strategy}-{repeat}.docx', 'equations': items[::-1]}
                (case / f'{strategy}-{repeat}.json').write_text(json.dumps(config, indent=2), encoding='utf-8')


def evidence(path: Path) -> tuple[ET.Element, Counter]:
    with zipfile.ZipFile(path) as archive:
        xml = ET.fromstring(archive.read('word/document.xml'))
        payloads = Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist() if name.startswith('word/embeddings/')
        )
    return xml, payloads


def sequence_results(xml: ET.Element) -> list[int]:
    results = []
    for paragraph in xml.iter(qn('w:p')):
        if not any('MACROBUTTON MTPlaceRef' in (element.text or '') for element in paragraph.iter(qn('w:instrText'))):
            continue
        # Word may serialize a nested sequence as fldSimple or as a complex field.
        flat_elements = list(paragraph.iter())
        marker_indices = [
            index for index, element in enumerate(flat_elements)
            if r'SEQ MTEqn \c' in ((element.text or '') + element.attrib.get(qn('w:instr'), ''))
        ]
        assert marker_indices, 'Missing displayed equation sequence.'
        value = next(
            element.text for element in flat_elements[marker_indices[-1] + 1:]
            if element.tag in {qn('w:instrText'), qn('w:t')} and (element.text or '').isdigit()
        )
        results.append(int(value))
    return results


def inspect(root: Path) -> None:
    rows = json.loads((root / 'measurements.json').read_text(encoding='utf-8-sig'))
    summaries = []
    for count in sorted({row['equations'] for row in rows}):
        case_rows = [row for row in rows if row['equations'] == count]
        _, seed_payloads = evidence(root / f'n-{count}' / 'display.docx')
        for row in case_rows:
            path = root / f'n-{count}' / f"{row['strategy']}-{row['repeat']}.docx"
            xml, payloads = evidence(path)
            assert sum(payloads.values()) == count
            assert sequence_results(xml) == list(range(1, count + 1))
            report = json.loads(path.with_suffix('.report.json').read_text(encoding='utf-8-sig'))
            assert len(report['equations']) == count
            assert all(item['progId'] == 'Equation.DSMT4' and item['nativeNumberFields'] > 0 for item in report['equations'])
            assert report['timings']['exportPdfMilliseconds'] == 0
            if row['strategy'] == 'reuse':
                assert payloads == seed_payloads, 'A reused OLE payload changed.'
                assert all(item['conversionMilliseconds'] == 0 for item in report['equations'])
            row['processingMilliseconds'] = report['timings']['processingMilliseconds']
            row['conversionMilliseconds'] = sum(item['conversionMilliseconds'] for item in report['equations'])
            row['numberingMilliseconds'] = sum(item['numberingMilliseconds'] for item in report['equations'])
            row['locateMilliseconds'] = report['timings']['locateAndValidateMilliseconds']
            row['fieldUpdateMilliseconds'] = report['timings']['numberFieldUpdateMilliseconds']
        medians = {
            strategy: round(statistics.median(row['wallMilliseconds'] for row in case_rows if row['strategy'] == strategy), 1)
            for strategy in ('recreate', 'reuse')
        }
        summaries.append({
            'equations': count,
            'repeatsPerStrategy': len(case_rows) // 2,
            'recreateMedianWallMilliseconds': medians['recreate'],
            'reuseMedianWallMilliseconds': medians['reuse'],
            'reuseWallReductionPercent': round(100 * (1 - medians['reuse'] / medians['recreate']), 1),
            'allOlePayloadsPreserved': True,
            'allSequenceResultsCorrect': True,
        })
    summary = {
        'comparison': 'Regenerate equations from LaTeX plus numbering vs number-existing; same content, no PDF.',
        'samples': rows,
        'summary': summaries,
    }
    (root / 'performance-summary.json').write_text(json.dumps(summary, indent=2), encoding='utf-8')
    print(json.dumps(summaries, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('action', choices=['prepare', 'inspect'])
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--counts', type=int, nargs='+', default=[10, 50])
    parser.add_argument('--repeats', type=int, default=3)
    parser.add_argument('--filler-paragraphs', type=int, default=500)
    args = parser.parse_args()
    if args.action == 'prepare':
        assert all(count > 0 for count in args.counts) and args.repeats > 0 and args.filler_paragraphs >= 0
        prepare(args.output.resolve(), args.counts, args.repeats, args.filler_paragraphs)
    else:
        inspect(args.output.resolve())
