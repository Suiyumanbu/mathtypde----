from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def write_manifest(root: Path, name: str, source: str, items: list[dict], pdf: bool = False) -> None:
    config = {'input': source, 'output': name + '.docx', 'equations': items}
    if pdf:
        config['pdf'] = name + '.pdf'
    (root / (name + '.json')).write_text(json.dumps(config, indent=2), encoding='utf-8')


def prepare(root: Path) -> None:
    write_manifest(root, 'already-numbered', 'renumbered.docx', [
        {'bookmark': 'mt_bookmark_display', 'operation': 'number-existing'},
    ], pdf=True)
    write_manifest(root, 'missing-bookmark', 'result.docx', [
        {'bookmark': 'mt_does_not_exist', 'operation': 'number-existing'},
    ], pdf=True)
    write_manifest(root, 'out-of-range', 'result.docx', [
        {'equationIndex': 999, 'operation': 'number-existing'},
    ], pdf=True)
    write_manifest(root, 'overlap', 'result.docx', [
        {'bookmark': 'mt_bookmark_display', 'operation': 'number-existing'},
        {'equationIndex': 4, 'operation': 'number-existing'},
    ], pdf=True)

    doc = Document(root / 'result.docx')
    # Keep the new formula after MathType's chapter-start field in the heading;
    # inserting before it intentionally starts a separate sequence segment.
    doc.paragraphs[1].insert_paragraph_before('[[EQ:mixed-new]]')
    doc.core_properties.author = 'MIXED_AUTHOR_SHOULD_NOT_APPEAR'
    numbered_paragraph = next(p for p in doc.paragraphs if 'MACROBUTTON MTPlaceRef' in p._p.xml)
    # Empty cached AUTHOR result beside a native equation number. Refreshing the
    # whole paragraph would populate it and violate unrelated-field preservation.
    run = numbered_paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.text = ' AUTHOR '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for element in (begin, instruction, separate, end):
        run._r.append(element)
    source = root / 'mixed-source.docx'
    doc.save(source)
    # Exercise restoration of an existing deferred-update property, not only deletion.
    custom_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/custom-properties'
    value_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'
    with zipfile.ZipFile(source) as archive:
        entries = {entry.filename: (entry, archive.read(entry)) for entry in archive.infolist()}
    custom = ET.fromstring(entries['docProps/custom.xml'][1])
    property_element = ET.SubElement(custom, f'{{{custom_ns}}}property', {
        'fmtid': '{D5CDD505-2E9C-101B-9397-08002B2CF9AE}',
        'pid': str(max(int(p.attrib['pid']) for p in custom) + 1),
        'name': 'MTDeferFieldUpdate',
    })
    ET.SubElement(property_element, f'{{{value_ns}}}lpwstr').text = '0'
    with zipfile.ZipFile(source, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, (entry, content) in entries.items():
            if name == 'docProps/custom.xml':
                content = ET.tostring(custom, encoding='utf-8', xml_declaration=True)
            archive.writestr(entry, content)
    write_manifest(root, 'mixed', 'mixed-source.docx', [
        {'equationIndex': 4, 'operation': 'number-existing'},
        {'anchor': '[[EQ:mixed-new]]', 'latex': 'F=ma', 'mode': 'right-numbered'},
    ])


def evidence(path: Path) -> tuple[str, str, Counter]:
    with zipfile.ZipFile(path) as archive:
        document = archive.read('word/document.xml').decode('utf-8')
        custom = archive.read('docProps/custom.xml').decode('utf-8')
        payloads = Counter(
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist() if name.startswith('word/embeddings/')
        )
    return document, custom, payloads


def inspect(root: Path) -> None:
    source_xml, _, source_payloads = evidence(root / 'mixed-source.docx')
    xml, custom, payloads = evidence(root / 'mixed.docx')
    report = json.loads((root / 'mixed.report.json').read_text(encoding='utf-8-sig'))
    assert sum(payloads.values()) == sum(source_payloads.values()) + 1
    assert not (source_payloads - payloads), 'A pre-existing OLE payload was changed.'
    assert xml.count('MACROBUTTON MTPlaceRef') == 4
    assert '[[EQ:mixed-new]]' not in xml
    assert 'MIXED_AUTHOR_SHOULD_NOT_APPEAR' not in xml, 'An adjacent AUTHOR field was updated.'
    assert 'UNCHANGED_FIELD' in xml
    assert ' AUTHOR ' in source_xml and ' AUTHOR ' in xml
    properties = ET.fromstring(custom)
    deferred = next(p for p in properties if p.attrib.get('name') == 'MTDeferFieldUpdate')
    assert next(iter(deferred)).text == '0', 'The original deferred-update property was not restored.'
    numbers = []
    for paragraph in re.findall(r'<w:p\b.*?</w:p>', xml):
        if 'MACROBUTTON MTPlaceRef' in paragraph:
            marker = paragraph.rfind(r'SEQ MTEqn \c')
            match = re.search(r'<w:instrText>(\d+)</w:instrText>', paragraph[marker:])
            assert marker >= 0 and match
            numbers.append(int(match.group(1)))
    assert numbers == [1, 2, 3, 4], numbers
    assert [item['source'] for item in report['equations']] == ['latex', 'existing']
    assert report['equations'][1]['conversionMilliseconds'] == 0
    assert report['visualReview'] == 'not-requested'
    assert report['timings']['exportPdfMilliseconds'] == 0
    assert not list(root.glob('.*.docx')), 'A failed job leaked staged DOCX files.'
    assert not list(root.glob('.*.pdf')), 'A failed job leaked staged PDFs.'
    print('Mixed insert/reuse, adjacent-field preservation, deferred-property restoration, and failure cleanup passed.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('action', choices=['prepare', 'inspect'])
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    {'prepare': prepare, 'inspect': inspect}[args.action](args.output.resolve())
